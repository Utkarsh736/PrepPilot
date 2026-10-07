"""Document parsing: PDF / DOCX / TXT / Markdown → plain text.

Uploads arrive as files (resumes are usually PDFs) or as pasted text (typical
for job descriptions). This module turns either into clean text plus some
light stats used for UI cards and logging.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

MAX_FILE_BYTES = 15 * 1024 * 1024  # 15 MB

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


@dataclass
class ParsedDocument:
    text: str
    pages: int  # page count for PDFs, 1 otherwise
    word_count: int
    parser: str
    warning: str = ""


class ParseError(ValueError):
    pass


def _normalize(text: str) -> str:
    # collapse Windows line endings & repeated blank lines
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ParseError("pypdf is not installed") from e

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:
                raise ParseError(
                    "This PDF is password-protected. Please remove the password and re-upload."
                ) from e
        pages_text = []
        for i, page in enumerate(reader.pages):
            try:
                pages_text.append(page.extract_text() or "")
            except Exception:
                pages_text.append("")
        text = "\n\n".join(pages_text)
    except ParseError:
        raise
    except Exception as e:
        raise ParseError(f"Could not read PDF: {e}") from e

    warning = ""
    clean = _normalize(text)
    if len(clean) < 40:
        warning = (
            "Very little text could be extracted — this PDF may be a scan/image. "
            "Try uploading a text-based PDF or paste the content instead."
        )
    return ParsedDocument(
        text=clean,
        pages=len(reader.pages),
        word_count=len(clean.split()),
        parser="pypdf",
        warning=warning,
    )


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        import docx  # type: ignore

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
        # also harvest tables (resumes often use tables for layout)
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        clean = _normalize("\n".join(parts))
        return ParsedDocument(
            text=clean,
            pages=1,
            word_count=len(clean.split()),
            parser="python-docx",
        )
    except Exception as e:
        raise ParseError(f"Could not read DOCX: {e}") from e


def parse_plain(data: bytes) -> ParsedDocument:
    try:
        text = data.decode("utf-8", errors="replace")
    except Exception as e:  # pragma: no cover
        raise ParseError(f"Could not decode text file: {e}") from e
    clean = _normalize(text)
    return ParsedDocument(
        text=clean,
        pages=1,
        word_count=len(clean.split()),
        parser="plaintext",
    )


def parse_upload(filename: str, data: bytes) -> ParsedDocument:
    """Dispatch on file extension. Raises ParseError on bad input."""
    if not data:
        raise ParseError("Empty file")
    if len(data) > MAX_FILE_BYTES:
        raise ParseError("File is larger than 15 MB")

    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return parse_pdf(data)
    if name.endswith(".docx"):
        return parse_docx(data)
    if name.endswith((".txt", ".md", ".markdown")):
        return parse_plain(data)

    # Unknown extension — try PDF magic bytes, else treat as plain text
    if data[:5] == b"%PDF-":
        return parse_pdf(data)
    if name.endswith(".doc"):
        raise ParseError(
            "Legacy .doc files are not supported — please save as .docx or PDF and re-upload."
        )
    return parse_plain(data)


def parse_pasted_text(text: str) -> ParsedDocument:
    clean = _normalize(text or "")
    if len(clean) < 20:
        raise ParseError("Pasted text is too short to be useful")
    return ParsedDocument(
        text=clean,
        pages=1,
        word_count=len(clean.split()),
        parser="pasted-text",
    )
